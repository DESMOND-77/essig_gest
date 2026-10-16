import datetime as dt

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from personnes.models import Etudiant, Inscription
from referentiel.models import (
    AnneeScolaire,
    Classe,
    Ecole,
    ModaliteFormation,
    Niveau,
    OptionFiliere,
)

# Feuilles ESSIG/ESSS/CFPA (§1 du mapping) ; mapping feuille -> Ecole.sigle,
# absent du fichier source, fourni en configuration (cf. db/IMPORT_MAPPING_scolarite.md §5).
SIGLE_ECOLE_PAR_FEUILLE = {"ESSIG": "ESSIG", "ESSS": "UPSS", "CFPA": "CFPA"}
LIGNE_PREMIERE_DONNEE = 4
MODALITE_DEFAUT = "FI"


class Command(BaseCommand):
    help = (
        "Import idempotent des inscriptions existantes (feuilles ESSIG/ESSS/CFPA) "
        "depuis le classeur legacy de scolarité (§16 Phase 10)."
    )

    def add_arguments(self, parser):
        parser.add_argument("fichier", help="Chemin du classeur .xlsx à importer.")
        parser.add_argument("--annee", default="2025-2026")
        parser.add_argument("--rapport-erreurs", default=None)

    def handle(self, *args, **opts):
        try:
            annee_scolaire = AnneeScolaire.objects.get(libelle=opts["annee"])
        except AnneeScolaire.DoesNotExist as exc:
            raise CommandError(
                f"AnneeScolaire '{opts['annee']}' introuvable."
            ) from exc

        classeur = openpyxl.load_workbook(opts["fichier"], data_only=True)
        erreurs = []
        total_traitees = total_importees = 0

        for nom_feuille, sigle_ecole in SIGLE_ECOLE_PAR_FEUILLE.items():
            if nom_feuille not in classeur.sheetnames:
                self.stdout.write(f"[{nom_feuille}] feuille absente, ignorée.")
                continue
            try:
                ecole = Ecole.objects.get(sigle=sigle_ecole)
            except Ecole.DoesNotExist:
                erreurs.append(
                    (nom_feuille, "-", f"École '{sigle_ecole}' introuvable en base.")
                )
                continue

            feuille = classeur[nom_feuille]
            with transaction.atomic():
                traitees, importees = self._importer_feuille(
                    feuille, nom_feuille, ecole, annee_scolaire, erreurs
                )
            total_traitees += traitees
            total_importees += importees
            self.stdout.write(
                f"[{nom_feuille}] {traitees} ligne(s) traitée(s), "
                f"{importees} inscription(s) importée(s)."
            )

        self.stdout.write(
            self.style.SUCCESS(
                f"Import terminé : {total_importees} inscription(s) importée(s) "
                f"sur {total_traitees} ligne(s) traitée(s), {len(erreurs)} erreur(s)."
            )
        )
        self._ecrire_rapport(opts["rapport_erreurs"], erreurs)

    def _importer_feuille(self, feuille, nom_feuille, ecole, annee_scolaire, erreurs):
        traitees = importees = 0
        for ligne in feuille.iter_rows(min_row=LIGNE_PREMIERE_DONNEE):
            valeurs = {c.column_letter: c.value for c in ligne}
            numero_ordre = valeurs.get("A")
            nom_prenoms = valeurs.get("B")
            if numero_ordre is None and not nom_prenoms:
                continue  # ligne de gabarit vide, pas une erreur (cf. mapping, constat).
            traitees += 1
            num_ligne = ligne[0].row
            try:
                if self._importer_ligne(
                    valeurs, num_ligne, ecole, annee_scolaire
                ):
                    importees += 1
            except LigneInvalide as exc:
                erreurs.append((nom_feuille, num_ligne, str(exc)))
                self.stdout.write(
                    self.style.WARNING(
                        f"[{nom_feuille}] ligne {num_ligne} ignorée : {exc}"
                    )
                )
        return traitees, importees

    def _importer_ligne(self, valeurs, num_ligne, ecole, annee_scolaire):
        matricule = (valeurs.get("AF") or valeurs.get("AE") or "")
        matricule = str(matricule).strip() if matricule else ""
        if not matricule:
            raise LigneInvalide(
                "étudiant sans matricule, ligne ignorée — à matriculer manuellement"
            )

        nom_prenoms = str(valeurs.get("B") or "").strip()
        nom, prenoms = self._separer_nom_prenoms(nom_prenoms, num_ligne)

        niveau_brut = str(valeurs.get("C") or "").strip()
        niveau = Niveau.objects.filter(ecole=ecole, libelle__iexact=niveau_brut).first()
        if niveau is None:
            niveau = Niveau.objects.filter(ecole=ecole, code__iexact=niveau_brut).first()
        if niveau is None:
            raise LigneInvalide(f'niveau "{niveau_brut}" inconnu pour l\'école {ecole}')

        filiere_option_brut = str(valeurs.get("D") or "").strip()
        option_filiere = self._resoudre_option_filiere(filiere_option_brut, ecole)
        if option_filiere is None:
            raise LigneInvalide(
                f'filière/option "{filiere_option_brut}" non résolue pour l\'école {ecole}'
            )

        modalite = self._resoudre_modalite(ecole)

        classe, _ = Classe.objects.get_or_create(
            annee_scolaire=annee_scolaire,
            niveau=niveau,
            option_filiere=option_filiere,
            modalite_formation=modalite,
            defaults={"libelle_auto": f"{niveau.code}-{option_filiere.sigle}"},
        )

        etudiant, _ = Etudiant.objects.update_or_create(
            matricule=matricule,
            defaults={"nom": nom, "prenoms": prenoms},
        )

        a_antecedent = (
            Inscription.objects.filter(etudiant=etudiant)
            .exclude(annee_scolaire=annee_scolaire)
            .exists()
        )
        type_inscription = (
            Inscription.Type.REINSCRIPTION if a_antecedent else Inscription.Type.INSCRIPTION
        )
        statut = Inscription.Statut.ANCIEN if a_antecedent else Inscription.Statut.NOUVEAU

        cout_brut = valeurs.get("AC")
        cout_negocie = None
        if isinstance(cout_brut, (int, float)) and cout_brut:
            cout_negocie = cout_brut

        Inscription.objects.update_or_create(
            etudiant=etudiant,
            annee_scolaire=annee_scolaire,
            defaults={
                "classe": classe,
                "ecole": ecole,
                "numero_ordre": int(valeurs.get("A") or 0),
                "type": type_inscription,
                "statut": statut,
                "cout_scolarite_negocie": cout_negocie,
                "date_inscription": annee_scolaire.date_debut,
                "etat": Inscription.Etat.ACTIVE,
            },
        )
        return True

    def _separer_nom_prenoms(self, nom_prenoms, num_ligne):
        mots = nom_prenoms.split()
        if len(mots) < 2:
            self.stdout.write(
                self.style.WARNING(
                    f"ligne {num_ligne} : nom/prénoms non séparables "
                    f'("{nom_prenoms}"), importés tel quel dans `nom`.'
                )
            )
            return nom_prenoms, ""
        return mots[0], " ".join(mots[1:])

    def _resoudre_option_filiere(self, brut, ecole):
        if "/" not in brut:
            return None
        sigle_filiere, _, sigle_option = (p.strip() for p in brut.partition("/"))
        return OptionFiliere.objects.filter(
            filiere__ecole=ecole,
            filiere__sigle__iexact=sigle_filiere,
            sigle__iexact=sigle_option,
        ).first()

    def _resoudre_modalite(self, ecole):
        return ModaliteFormation.objects.filter(code=MODALITE_DEFAUT).first()

    def _ecrire_rapport(self, chemin, erreurs):
        if not erreurs and chemin is None:
            return
        if chemin is None:
            horodatage = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
            chemin = f"rapport_import_scolarite_{horodatage}.txt"
        with open(chemin, "w", encoding="utf-8") as fichier:
            for feuille, ligne, raison in erreurs:
                fichier.write(f"{feuille}\tligne {ligne}\t{raison}\n")
        self.stdout.write(f"Rapport d'erreurs écrit : {chemin}")


class LigneInvalide(Exception):
    pass
