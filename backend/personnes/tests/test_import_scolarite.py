import openpyxl
import pytest
from django.core.management import call_command

from personnes.models import Etudiant, Inscription
from referentiel.tests.factories import (
    AnneeScolaireFactory,
    EcoleFactory,
    FiliereFactory,
    ModaliteFormationFactory,
    NiveauFactory,
    OptionFiliereFactory,
)

pytestmark = pytest.mark.django_db


def _creer_classeur(tmp_path, niveau, option_filiere):
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "ESSIG"
    for ligne in range(1, 4):
        feuille.cell(row=ligne, column=1, value="en-tête")
    feuille.cell(row=4, column=1, value=1)  # A: N° ordre
    feuille.cell(row=4, column=2, value="DIALLO Mamadou")  # B: NOMS ET PRENOMS
    feuille.cell(row=4, column=3, value=niveau.libelle)  # C: NIVEAU
    feuille.cell(
        row=4, column=4, value=f"{option_filiere.filiere.sigle}/{option_filiere.sigle}"
    )  # D
    feuille.cell(row=4, column=29, value=150000)  # AC: Coût de la scolarité
    feuille.cell(row=4, column=32, value="ESSIG-0001")  # AF: N° MATRICULE
    chemin = tmp_path / "test_import.xlsx"
    classeur.save(chemin)
    return chemin


@pytest.fixture
def jeu_de_test():
    ecole = EcoleFactory(sigle="ESSIG")
    annee = AnneeScolaireFactory(libelle="2025-2026")
    niveau = NiveauFactory(ecole=ecole, libelle="L1")
    filiere = FiliereFactory(ecole=ecole, sigle="GACO")
    option = OptionFiliereFactory(filiere=filiere, sigle="CI")
    modalite = ModaliteFormationFactory(code="FI")
    return ecole, annee, niveau, option, modalite


def test_import_scolarite_cree_etudiant_et_inscription(tmp_path, jeu_de_test):
    _, annee, niveau, option, _ = jeu_de_test
    chemin = _creer_classeur(tmp_path, niveau, option)

    call_command("import_scolarite", str(chemin), "--annee", annee.libelle)

    assert Etudiant.objects.filter(matricule="ESSIG-0001").count() == 1
    assert Inscription.objects.filter(etudiant__matricule="ESSIG-0001").count() == 1
    inscription = Inscription.objects.get(etudiant__matricule="ESSIG-0001")
    assert inscription.cout_scolarite_negocie == 150000


def test_import_scolarite_est_idempotent(tmp_path, jeu_de_test):
    _, annee, niveau, option, _ = jeu_de_test
    chemin = _creer_classeur(tmp_path, niveau, option)

    call_command("import_scolarite", str(chemin), "--annee", annee.libelle)
    call_command("import_scolarite", str(chemin), "--annee", annee.libelle)

    assert Etudiant.objects.filter(matricule="ESSIG-0001").count() == 1
    assert Inscription.objects.filter(etudiant__matricule="ESSIG-0001").count() == 1


def test_import_scolarite_journalise_ligne_sans_matricule(tmp_path, jeu_de_test):
    _, annee, niveau, option, _ = jeu_de_test
    chemin = _creer_classeur(tmp_path, niveau, option)
    classeur = openpyxl.load_workbook(chemin)
    classeur["ESSIG"].cell(row=4, column=32).value = None
    classeur.save(chemin)
    rapport = tmp_path / "rapport.txt"

    call_command(
        "import_scolarite",
        str(chemin),
        "--annee",
        annee.libelle,
        "--rapport-erreurs",
        str(rapport),
    )

    assert Etudiant.objects.count() == 0
    assert rapport.exists()
    assert "matricule" in rapport.read_text(encoding="utf-8")
