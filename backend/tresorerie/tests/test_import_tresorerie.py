from datetime import date
from io import StringIO

import openpyxl
import pytest
from django.core.management import call_command

from referentiel.models import AnneeScolaire
from tresorerie.models import Depense, Mouvement


def construire_classeur(tmp_path):
    classeur = openpyxl.Workbook()

    feuille_depenses = classeur.active
    feuille_depenses.title = "DEPENSES"
    feuille_depenses["A5"], feuille_depenses["B5"], feuille_depenses["C5"] = "DATES", "DESIGNATIONS", "SOMMES"
    feuille_depenses["A6"], feuille_depenses["B6"], feuille_depenses["C6"] = date(2025, 10, 5), "Fournitures", 15000
    feuille_depenses["A7"], feuille_depenses["B7"], feuille_depenses["C7"] = "Total", None, 15000

    feuille_versements = classeur.create_sheet("VERSEMENTS ")
    feuille_versements["B5"], feuille_versements["C5"] = "DATES", "SOMMES"
    feuille_versements["F5"], feuille_versements["G5"] = "DATES", "SOMMES"
    feuille_versements["B6"], feuille_versements["C6"] = date(2025, 10, 6), 50000
    feuille_versements["F6"], feuille_versements["G6"] = date(2025, 10, 7), 10000

    chemin = tmp_path / "test_import.xlsx"
    classeur.save(chemin)
    return chemin


@pytest.mark.django_db
def test_import_tresorerie_est_idempotent(tmp_path):
    AnneeScolaire.objects.create(
        libelle="2025-2026", date_debut="2025-10-01", date_fin="2026-08-31",
    )
    chemin = construire_classeur(tmp_path)

    args = [
        str(chemin),
        "--annee-scolaire", "2025-2026",
        "--compte-caisse", "CAISSE-TEST",
        "--compte-banque", "BANQUE-TEST",
    ]

    call_command("import_tresorerie", *args, stdout=StringIO())
    assert Depense.objects.count() == 1
    assert Mouvement.objects.count() == 3  # 1 dépense + 1 versement + 1 retrait

    call_command("import_tresorerie", *args, stdout=StringIO())
    assert Depense.objects.count() == 1
    assert Mouvement.objects.count() == 3


@pytest.mark.django_db
def test_import_tresorerie_journalise_les_lignes_invalides(tmp_path):
    AnneeScolaire.objects.create(
        libelle="2025-2026", date_debut="2025-10-01", date_fin="2026-08-31",
    )
    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.title = "DEPENSES"
    feuille["A5"], feuille["B5"], feuille["C5"] = "DATES", "DESIGNATIONS", "SOMMES"
    feuille["A6"], feuille["B6"], feuille["C6"] = None, "Désignation sans date", 5000
    chemin = tmp_path / "test_erreurs.xlsx"
    classeur.save(chemin)

    rapport = tmp_path / "rapport.txt"
    call_command(
        "import_tresorerie", str(chemin),
        "--annee-scolaire", "2025-2026",
        "--compte-caisse", "CAISSE-TEST",
        "--compte-banque", "BANQUE-TEST",
        "--rapport-erreurs", str(rapport),
        stdout=StringIO(),
    )

    assert Depense.objects.count() == 0
    assert rapport.exists()
    assert "date manquante" in rapport.read_text(encoding="utf-8")
