"""Import legacy du groupe `enseignants` (§16 Phase 10) depuis `Enseignants.xlsx`.

Mapping de référence : `db/IMPORT_MAPPING_enseignants.md`. Import d'identité
pure (nom/prénom/téléphone/email) ; `specialite`, `statut`, `taux_horaire`
restent à saisir/importer séparément (cf. mapping §1/§5).
"""
from datetime import datetime
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from rh_paie.models import Enseignant

LIGNE_PREMIERE_DONNEE = 2


class Command(BaseCommand):
    help = (
        "Import idempotent des enseignants existants (feuille unique "
        "NOM/PRENOM/TELEPHONE/EMAIL) depuis un classeur legacy (§16 Phase 10)."
    )

    def add_arguments(self, parser):
        parser.add_argument("fichier", help="Chemin du classeur .xlsx à importer.")
        parser.add_argument("--rapport-erreurs", default=None)

    def handle(self, *args, **opts):
        chemin = Path(opts["fichier"])
        if not chemin.exists():
            raise CommandError(f"Fichier introuvable : {chemin}")

        classeur = openpyxl.load_workbook(chemin, data_only=True)
        feuille = classeur[classeur.sheetnames[0]]

        erreurs = []
        traitees = importees = 0

        with transaction.atomic():
            for ligne in feuille.iter_rows(min_row=LIGNE_PREMIERE_DONNEE):
                nom_brut, prenom_brut, tel_brut, email_brut = (
                    c.value for c in ligne[:4]
                )
                if nom_brut is None and prenom_brut is None:
                    continue  # ligne vide, pas une erreur (cf. mapping §4.9).
                num_ligne = ligne[0].row
                traitees += 1
                try:
                    self._importer_ligne(nom_brut, prenom_brut, tel_brut, email_brut)
                    importees += 1
                except LigneInvalide as exc:
                    erreurs.append((num_ligne, str(exc)))
                    self.stdout.write(
                        self.style.WARNING(f"ligne {num_ligne} ignorée : {exc}")
                    )

        self.stdout.write(
            self.style.SUCCESS(
                f"Import terminé : {importees} enseignant(s) importé(s) sur "
                f"{traitees} ligne(s) traitée(s), {len(erreurs)} erreur(s)."
            )
        )
        self._ecrire_rapport(opts["rapport_erreurs"], erreurs)

    def _importer_ligne(self, nom_brut, prenom_brut, tel_brut, email_brut):
        nom = str(nom_brut or "").strip()
        prenom = str(prenom_brut or "").strip().replace("\\'", "'")
        if not nom or not prenom:
            raise LigneInvalide("NOM ou PRENOM vide")

        telephone = self._normaliser_telephone(tel_brut)
        email = self._normaliser_email(email_brut, nom, prenom)

        Enseignant.objects.update_or_create(
            nom__iexact=nom,
            prenom__iexact=prenom,
            defaults={"nom": nom, "prenom": prenom, "telephone": telephone, "email": email},
        )

    def _normaliser_telephone(self, tel_brut):
        if tel_brut is None:
            return None
        if isinstance(tel_brut, (int, float)):
            self.stdout.write(
                self.style.WARNING(
                    f"TELEPHONE numérique converti en texte ({tel_brut}), "
                    "vérifier un éventuel '+'/'0' de tête perdu."
                )
            )
            return str(int(tel_brut))
        return str(tel_brut).strip()

    def _normaliser_email(self, email_brut, nom, prenom):
        if not email_brut:
            return None
        email = str(email_brut).strip().lower()
        if "@" not in email:
            self.stdout.write(
                self.style.WARNING(f'EMAIL au format invalide ignoré : "{email}"')
            )
            return None
        conflit = (
            Enseignant.objects.filter(email=email)
            .exclude(nom__iexact=nom, prenom__iexact=prenom)
            .exists()
        )
        if conflit:
            self.stdout.write(
                self.style.WARNING(
                    f'email déjà utilisé par un autre enseignant : "{email}" '
                    f"({nom} {prenom} importé sans email)"
                )
            )
            return None
        return email

    def _ecrire_rapport(self, chemin, erreurs):
        if not erreurs and chemin is None:
            return
        if chemin is None:
            horodatage = datetime.now().strftime("%Y%m%d_%H%M%S")
            chemin = f"rapport_import_enseignants_{horodatage}.txt"
        with open(chemin, "w", encoding="utf-8") as fichier:
            for num_ligne, raison in erreurs:
                fichier.write(f"ligne {num_ligne}\t{raison}\n")
        self.stdout.write(f"Rapport d'erreurs écrit : {chemin}")


class LigneInvalide(Exception):
    pass
